
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

import pymupdf
from docx import Document

from grading_models import RubricCreate


BASE_DIR = Path(__file__).resolve().parent
GRADING_DIR = BASE_DIR / "grading_data"

RUBRIC_DIR = GRADING_DIR / "rubrics"
ESSAY_DIR = GRADING_DIR / "essays"
GRADE_DIR = GRADING_DIR / "grades"

for directory in (RUBRIC_DIR, ESSAY_DIR, GRADE_DIR):
    directory.mkdir(parents=True, exist_ok=True)

MAX_ESSAY_CHARACTERS = 18000


def now_utc():
    return datetime.now(timezone.utc).isoformat()


def valid_id(value: str):
    return str(UUID(value))


def save_json(path: Path, data: dict):
    """
    Write JSON atomically to avoid partial files.
    """
    temporary = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            suffix=".tmp",
            delete=False
        ) as file:
            temporary = Path(file.name)
            json.dump(
                data,
                file,
                ensure_ascii=False,
                indent=2
            )

        os.replace(temporary, path)

    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def read_json(path: Path):
    if not path.exists():
        return None

    with path.open(encoding="utf-8") as file:
        return json.load(file)


def save_rubric(rubric: RubricCreate):
    rubric_id = str(uuid4())

    data = {
        "rubric_id": rubric_id,
        "created_at": now_utc(),
        **rubric.model_dump()
    }

    save_json(
        RUBRIC_DIR / f"{rubric_id}.json",
        data
    )

    return data


def get_rubric(rubric_id: str):
    rubric_id = valid_id(rubric_id)

    return read_json(
        RUBRIC_DIR / f"{rubric_id}.json"
    )


def list_rubrics():
    result = []

    for path in sorted(RUBRIC_DIR.glob("*.json")):
        data = read_json(path)

        result.append({
            "rubric_id": data["rubric_id"],
            "title": data["title"],
            "created_at": data["created_at"],
            "total_criteria": len(data["criteria"])
        })

    return result


def save_essay(title: str, text: str, filename=None):
    text = text.strip()

    if not text:
        raise ValueError("Essay is empty.")

    if len(text) > MAX_ESSAY_CHARACTERS:
        raise ValueError(
            f"Essay exceeds {MAX_ESSAY_CHARACTERS} characters."
        )

    essay_id = str(uuid4())

    data = {
        "essay_id": essay_id,
        "title": title.strip()[:200],
        "filename": filename,
        "text": text,
        "character_count": len(text),
        "created_at": now_utc()
    }

    save_json(
        ESSAY_DIR / f"{essay_id}.json",
        data
    )

    return data


def get_essay(essay_id: str):
    essay_id = valid_id(essay_id)

    return read_json(
        ESSAY_DIR / f"{essay_id}.json"
    )


def save_grade(report: dict):
    grade_id = str(uuid4())

    data = {
        "grade_id": grade_id,
        "created_at": now_utc(),
        **report
    }

    save_json(
        GRADE_DIR / f"{grade_id}.json",
        data
    )

    return data


def get_grade(grade_id: str):
    grade_id = valid_id(grade_id)

    return read_json(
        GRADE_DIR / f"{grade_id}.json"
    )


def extract_essay_text(
    filename: str,
    content: bytes
) -> str:
    """
    Extract text from TXT, PDF, or DOCX.

    Scanned PDFs are not OCR-processed.
    """

    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError(
                "TXT files must use UTF-8 encoding."
            )

    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF header.")

        try:
            with pymupdf.open(
                stream=content,
                filetype="pdf"
            ) as pdf:
                if pdf.needs_pass:
                    raise ValueError(
                        "Password-protected PDFs are unsupported."
                    )

                return "\n\n".join(
                    page.get_text("text", sort=True)
                    for page in pdf
                )
        except pymupdf.FileDataError:
            raise ValueError("Invalid PDF file.")

    if extension == ".docx":
        import io

        try:
            document = Document(io.BytesIO(content))
        except Exception:
            raise ValueError("Invalid DOCX file.")

        paragraphs = [
            p.text
            for p in document.paragraphs
            if p.text.strip()
        ]

        # Include text from ordinary tables.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n\n".join(paragraphs)

    raise ValueError(
        "Supported essay formats: TXT, PDF, DOCX."
    )
